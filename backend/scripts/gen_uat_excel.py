"""Generate RadioLife_UAT_TestCases.xlsx — a ready-to-run UAT matrix."""
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter

OUT = "/app/RadioLife_UAT_TestCases.xlsx"

# TC = (Module, Feature, Priority, Platform, Type, Preconditions, Steps, TestData, Expected)
CASES = [
    # 1. AUTH & ONBOARDING
    ("Authentication", "Login", "P0", "Both", "Functional",
     "Registered user; on Login screen",
     "1. Enter email & password\n2. Tap/Click Sign in",
     "demo@bigfm.dubai / DemoPass123!",
     "User is authenticated and lands on Home; session persists on reload"),
    ("Authentication", "Login - wrong password", "P1", "Both", "Negative",
     "On Login screen",
     "1. Enter valid email + wrong password\n2. Submit",
     "demo@bigfm.dubai / wrongpass",
     "Friendly error shown; NOT logged in; no crash; no 'clear cache' prompt"),
    ("Authentication", "Onboarding gate", "P1", "Both", "Functional",
     "New user with onboarding_complete=false",
     "1. Log in as a fresh user\n2. Observe redirect",
     "new signup",
     "User is redirected to onboarding, not Home, until onboarding completes"),
    ("Authentication", "Logout", "P2", "Both", "Functional",
     "Logged in",
     "1. Open profile/menu\n2. Tap Logout",
     "-",
     "Session cleared; redirected to Login; protected routes no longer accessible"),

    # 2. HOME FEED / SHELL
    ("Home / Shell", "Home loads with shelves", "P0", "Web", "Smoke",
     "Logged in",
     "1. Open Home (/)\n2. Wait for load",
     "-",
     "Header, filter chips, and shelves render; no blank screen; no console errors"),
    ("Home / Shell", "Sidebar chip counts", "P1", "Web", "Functional",
     "Logged in",
     "1. Observe left sidebar / filter chips\n2. Compare counts to content",
     "-",
     "Chips show live counts (Radio, Shop cats, News, Games); zero-count items hidden"),
    ("Home / Shell", "Radio chip switches feed", "P0", "Web", "Functional",
     "On Home, chip=All",
     "1. Click 'Radio' chip\n2. Observe feed",
     "-",
     "Feed switches to a station grid (home-station-grid); visibly different from default feed"),
    ("Home / Shell", "Empty-state doctrine", "P1", "Both", "Functional",
     "A shelf/category has no content",
     "1. View Home default feed\n2. Look for empty shelves",
     "-",
     "Shelves with no content are hidden entirely; no 'nothing yet' placeholder in default feed"),

    # 3. RADIO & PLAYBACK
    ("Radio & Playback", "Play an AI channel", "P0", "Both", "Functional",
     "Logged in",
     "1. Open a station tile in 'Your stations'\n2. Tap play",
     "any channel",
     "Audio starts; persistent player bar shows station + now-playing segment"),
    ("Radio & Playback", "Play imported world station", "P0", "Both", "Functional",
     "On Radio browse / all stations",
     "1. Search/browse to a world station\n2. Tap play",
     "e.g. Classic FM UK",
     "Stream plays; station detail shows tags/country/language"),
    ("Radio & Playback", "Persistent player continue", "P1", "Both", "Functional",
     "A station is playing",
     "1. Navigate to another tab/page\n2. Observe player",
     "-",
     "Playback continues; player bar persists; Continue-listening shows the station"),
    ("Radio & Playback", "Search stations", "P1", "Both", "Functional",
     "On stations/search",
     "1. Type a query\n2. Review results",
     "'jazz'",
     "Relevant stations returned; each shows artwork (never blank)"),

    # 4. STATION ARTWORK (NEW)
    ("Station Artwork", "Every station has art", "P0", "Web", "Functional",
     "On Radio browse",
     "1. Open browse list\n2. Inspect each station tile",
     "-",
     "Every tile shows an image: real logo, AI art, or genre-coloured initials placeholder — never blank"),
    ("Station Artwork", "On-demand AI on first open", "P1", "Web", "Functional",
     "Open a placeholder-only imported station",
     "1. Open station detail (/radio/station/imp_...)\n2. Wait ~20-40s, refresh",
     "a station with placeholder art",
     "AI cover is generated from name/genre and appears on refresh; art_source becomes 'ai'"),
    ("Station Artwork", "Owner artwork manager visible", "P1", "Web", "Functional",
     "Logged in as admin OR station owner",
     "1. Open a station detail page\n2. Look for owner tools",
     "admin@bigfm.dubai / AdminPass123!",
     "'Station artwork · owner tools' panel shows Generate / Upload / Reset controls"),
    ("Station Artwork", "Owner controls hidden for non-owner", "P0", "Web", "Negative",
     "Logged in as a normal (non-owner) user",
     "1. Open the same station detail page",
     "demo@bigfm.dubai / DemoPass123!",
     "Owner artwork panel is NOT shown; generate/upload/clear API returns 403 if called"),
    ("Station Artwork", "Owner generates AI art", "P1", "Web", "Functional",
     "Logged in as admin/owner on station detail",
     "1. Click 'Generate AI artwork'\n2. Wait for completion toast",
     "-",
     "New AI cover replaces placeholder; success toast; art persists on reload"),
    ("Station Artwork", "Owner uploads logo", "P1", "Web", "Functional",
     "Logged in as admin/owner",
     "1. Click 'Upload logo'\n2. Choose a JPG/PNG/WebP < 5MB",
     "sample logo",
     "Uploaded image becomes the station art (art_source='owner'); >5MB or wrong type rejected with a clear error"),
    ("Station Artwork", "Owner resets art", "P2", "Web", "Functional",
     "Owner station has custom art",
     "1. Click 'Reset'",
     "-",
     "Art reverts to default/placeholder; can be regenerated again"),
    ("Station Artwork", "Flagship shelf self-heal", "P1", "Web", "Functional",
     "Flagship missing art (fresh deploy)",
     "1. Load Home\n2. Wait ~18s and refresh",
     "-",
     "Flagship tiles without art auto-generate AI covers and appear; no blank flagship tiles"),

    # 5. CLIPS & SPONSORED / TAP-TO-CLAIM (NEW)
    ("Clips", "Clips shelf renders", "P1", "Both", "Functional",
     "Logged in",
     "1. Scroll to 'Clips' shelf on Home",
     "-",
     "Vertical 9:16 clip cards render with artwork; shelf hidden if empty"),
    ("Clips", "Sponsored clip is first + badged", "P0", "Both", "Functional",
     "Active advertiser campaign exists",
     "1. View Clips shelf",
     "-",
     "First clip shows a 'SPONSORED' badge + branded cover image (not a plain gradient)"),
    ("Clips", "Tap-to-Claim issues code", "P0", "Web", "Functional",
     "Logged in; sponsored clip present",
     "1. Tap the sponsored clip\n2. In the sheet, tap 'Claim this offer'",
     "-",
     "Claim sheet opens; on claim a redemption code (RL-XXXXXX) + success state shown"),
    ("Clips", "Claim is idempotent", "P1", "Web", "Functional",
     "User already claimed this offer",
     "1. Tap sponsored clip again\n2. Claim",
     "same user + campaign",
     "Same code returned; 'Already claimed' state; no duplicate credit"),
    ("Clips", "Station clip opens station", "P2", "Both", "Functional",
     "Clips shelf has station clips",
     "1. Tap a non-sponsored clip",
     "-",
     "Navigates to the station / radio player"),

    # 6. SHOP & EAT
    ("Shop & Eat", "Category chips filter venues", "P0", "Web", "Functional",
     "On Home with venues",
     "1. Click 'Eat' chip\n2. Then 'Salons', 'Fitness', 'Shops'",
     "-",
     "Venue shelf filters to only that category each time; counts match"),
    ("Shop & Eat", "Sponsored venue marked", "P0", "Web", "Functional",
     "A venue has an active promotion",
     "1. View venue shelf",
     "-",
     "First card of a promoted set shows a 'Sponsored' pill; ordering is sponsored-first"),
    ("Shop & Eat", "Market-scoped venues", "P1", "Both", "Functional",
     "User geo resolves to Dallas",
     "1. Open Home venue shelf",
     "Dallas-IP user",
     "Shows Dallas venues (Smoke & Oak BBQ, etc.); Dubai user sees Dubai venues; never empty"),
    ("Shop & Eat", "Sell CTA at bottom", "P2", "Web", "Functional",
     "On /shop",
     "1. Scroll to bottom of Shop page",
     "-",
     "'Sell on RadioLife' / 'Merchant Dashboard' banner appears at the BOTTOM, not between categories and content"),
    ("Shop & Eat", "Order-by-chat floating capsule", "P2", "Web", "Functional",
     "On /shop",
     "1. Observe right side of screen\n2. Tap the capsule",
     "-",
     "'Order by chat' floating capsule pinned bottom-right; opens the AI concierge"),
    ("Shop & Eat", "Quick-Buy sheet", "P2", "Web", "Functional",
     "On a shopping sub-category",
     "1. Tap a product\n2. Use Quick-Buy",
     "-",
     "Quick-Buy sheet opens; memberships/discounts applied; add to cart works"),

    # 7. BIGPOINTS & WALLET
    ("BigPoints & Wallet", "Currency shows BigPoints", "P1", "Both", "Regression",
     "Anywhere currency is shown",
     "1. Check wallet, earn labels, offers",
     "-",
     "All currency reads 'BigPoints' — NO 'BigCoin' or 'BC' anywhere"),
    ("BigPoints & Wallet", "Wallet balance", "P1", "Both", "Functional",
     "Logged in",
     "1. Open Wallet",
     "-",
     "Balance + tier shown; matches points pill in header"),
    ("BigPoints & Wallet", "Earn from action", "P1", "Both", "Functional",
     "Logged in",
     "1. Complete an earn action (e.g. claim / play)\n2. Check wallet",
     "-",
     "Balance increases by the expected amount; transaction logged"),

    # 8. ON-AIR OFFERS
    ("On-Air Offers", "Offer card while playing", "P1", "Web", "Functional",
     "A station with an active offer is playing",
     "1. Play the station\n2. Observe card above player",
     "-",
     "On-air offer card shows the relevant merchant offer synced to the station"),
    ("On-Air Offers", "Claim from on-air card", "P2", "Web", "Functional",
     "On-air offer card visible",
     "1. Tap the offer's claim CTA",
     "-",
     "Routes to the offer/claim; no dead-end"),

    # 9. GAMES
    ("Games - Predict & Win", "Predictions list loads", "P1", "Web", "Functional",
     "Logged in",
     "1. Open Predict & Win (/predict-win)",
     "-",
     "Curated prediction markets load; categories filterable; no junk/unapproved questions"),
    ("Games - Predict & Win", "Make a pick", "P1", "Web", "Functional",
     "On a prediction",
     "1. Select an option\n2. Submit",
     "-",
     "Pick locked in; status shows 'Locked in'; reflected in My Picks"),
    ("Games - Trivia", "Trivia categories load", "P1", "Web", "Functional",
     "Logged in",
     "1. Open Trivia (/trivia)",
     "-",
     "City-aware categories shown (e.g. DFW/your city); header shows resolved market"),
    ("Games", "Spin & Win / Play hub", "P2", "Both", "Functional",
     "Logged in",
     "1. Open Play/Games",
     "-",
     "Live games listed; can start a game; earns BigPoints"),

    # 10. NEWS
    ("News", "City-aware headlines", "P1", "Both", "Functional",
     "Geo resolved",
     "1. Open Home news shelf / News tab",
     "Dallas-IP user",
     "Headlines request passes the detected market; feed reflects the city where possible"),
    ("News", "Breaking banner", "P2", "Both", "Functional",
     "A breaking item exists",
     "1. Observe top of Home / News",
     "-",
     "Breaking banner shown; tapping opens the story/News"),

    # 11. ADS
    ("Ads", "Leaderboard + in-feed banners", "P1", "Web", "Functional",
     "Non-premium user",
     "1. Load Home\n2. Observe ad slots",
     "-",
     "Banner slots render from active campaigns; premium users see fewer/no ads"),
    ("Ads", "Impression/click tracking", "P2", "Web", "Functional",
     "Sponsored station/clip shown",
     "1. View then click a sponsored item",
     "-",
     "Impression fired on view; click fired on tap (verify via network/logs)"),

    # 12. DESIGN / THEME
    ("Design / Theme", "Dark/light toggle", "P1", "Web", "Functional",
     "Logged in",
     "1. Toggle theme (theme-toggle)\n2. Navigate key pages",
     "-",
     "Theme flips app-wide; persists; no washed-out or unreadable elements"),
    ("Design / Theme", "Predict & Win light theme", "P0", "Web", "Functional",
     "Theme = light",
     "1. Open /predict-win",
     "-",
     "Background/text/cards/chips all legible in light; no black-on-black or white-on-white"),
    ("Design / Theme", "Trivia light theme", "P1", "Web", "Functional",
     "Theme = light",
     "1. Open /trivia",
     "-",
     "Category cards + sidebar legible in light and dark"),

    # 13. LOCATION
    ("Location", "Single geo source", "P1", "Both", "Functional",
     "Fresh session",
     "1. Load app\n2. Check header city, news, venues, trivia city",
     "-",
     "All reflect ONE detected city; no hardcoded Orlando/NY/Dubai mismatch"),

    # 14. MOBILE APP
    ("Mobile App", "Home dark YouTube shell", "P0", "Mobile", "Smoke",
     "Logged in on Expo app",
     "1. Open Home tab",
     "-",
     "Dark themed Home; header with menu + points; sticky filter chips; no blank screen"),
    ("Mobile App", "Swipeable shelves w/ peek", "P1", "Mobile", "Functional",
     "On mobile Home",
     "1. Swipe each shelf horizontally",
     "-",
     "Shelves scroll horizontally; next card peeks (half-card); smooth 60fps"),
    ("Mobile App", "Slide-out drawer", "P1", "Mobile", "Functional",
     "On mobile Home",
     "1. Tap the menu (hamburger)\n2. Tap a drawer item",
     "-",
     "Drawer slides in with Primary + Discover nav; item navigates and drawer closes"),
    ("Mobile App", "Bottom tab bar", "P1", "Mobile", "Functional",
     "On any tab",
     "1. Tap between tabs",
     "-",
     "Dark tab bar; selected tab clearly differs in colour from unselected; correct screen loads"),
    ("Mobile App", "Sponsored clip claim (mobile)", "P1", "Mobile", "Functional",
     "Sponsored clip on mobile Home",
     "1. Tap the sponsored clip",
     "-",
     "Claim fires; an inline banner shows the redemption code; dismissable"),
    ("Mobile App", "Pull to refresh", "P2", "Mobile", "Functional",
     "On mobile Home",
     "1. Pull down to refresh",
     "-",
     "Feed reloads; spinner uses brand colour"),
    ("Mobile App", "Safe areas", "P2", "Mobile", "Functional",
     "Device with notch",
     "1. Observe header + tab bar",
     "-",
     "Content respects safe areas; nothing clipped by notch/home indicator"),

    # 15. MERCHANT
    ("Merchant", "Merchant dashboard access", "P1", "Web", "Functional",
     "Logged in as merchant",
     "1. Open Merchant Dashboard",
     "test@bigfm.dubai / TestUser123!",
     "Dashboard loads; can manage menu/products/promotions/memberships"),
    ("Merchant", "Create a promotion", "P2", "Web", "Functional",
     "Merchant with a business",
     "1. Create/boost a promotion",
     "-",
     "Business becomes sponsored; appears sponsored-first on the venue shelf"),

    # 16. ADMIN
    ("Admin", "Admin login + panel", "P1", "Web", "Functional",
     "Admin account",
     "1. Log in as admin\n2. Open admin routes",
     "admin@bigfm.dubai / AdminPass123!",
     "Admin panel accessible; non-admins are blocked from admin routes"),
    ("Admin", "Approvals / sponsorships", "P2", "Web", "Functional",
     "Pending items exist",
     "1. Review approvals / trivia & station sponsorships",
     "-",
     "Admin can approve/reject; changes reflect in consumer app"),

    # 17. CROSS-PLATFORM PARITY
    ("Cross-platform", "Same data web vs mobile", "P1", "Both", "Functional",
     "Same user on both",
     "1. Compare Home shelves on web and mobile",
     "demo@bigfm.dubai",
     "Stations, clips, venues, points match (served by one /api); no platform-only divergence"),

    # 18. NON-FUNCTIONAL
    ("Non-functional", "Home load performance", "P2", "Both", "Performance",
     "Normal network",
     "1. Cold load Home\n2. Measure time to interactive",
     "-",
     "Home usable within a few seconds; skeletons/loaders shown while fetching"),
    ("Non-functional", "Graceful API failure", "P1", "Both", "Negative",
     "Simulate a failed endpoint",
     "1. Block one shelf's API\n2. Load Home",
     "-",
     "That shelf is hidden or shows a friendly error+retry; rest of Home still works; no crash"),
    ("Non-functional", "Broken image fallback", "P2", "Both", "Negative",
     "A station has a dead image URL",
     "1. Open that station",
     "-",
     "Falls back to placeholder/initials; layout not broken"),
    ("Non-functional", "Deep link / refresh", "P2", "Web", "Functional",
     "On a deep page",
     "1. Refresh the browser on /shop, /trivia, /predict-win",
     "-",
     "Page reloads correctly (no 404); auth preserved"),
]

HEADERS = ["TC ID", "Module", "Feature", "Priority", "Platform", "Type",
           "Preconditions", "Test Steps", "Test Data", "Expected Result",
           "Actual Result", "Status", "Tester", "Date", "Notes / Defect ID"]

# ---- styling helpers ----
BRAND = "0A0A0F"
ACCENT = "22D3EE"
HDR_FILL = PatternFill("solid", fgColor=BRAND)
HDR_FONT = Font(bold=True, color="FFFFFF", size=11)
WRAP = Alignment(wrap_text=True, vertical="top")
CENTER = Alignment(horizontal="center", vertical="top")
THIN = Side(style="thin", color="D9D9D9")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
PRIO_FILL = {"P0": "F8D7DA", "P1": "FFF3CD", "P2": "D1E7DD"}

wb = openpyxl.Workbook()

# ---------- Sheet 1: Instructions ----------
ws0 = wb.active
ws0.title = "Instructions"
ws0.sheet_view.showGridLines = False
ws0.column_dimensions["A"].width = 3
ws0.column_dimensions["B"].width = 110
lines = [
    ("RadioLife.app — UAT & Test Cases", "title"),
    ("", ""),
    ("Environments", "h"),
    ("• PRODUCTION (live): https://radiolife.app — test the real, deployed app here.", "p"),
    ("• PREVIEW (dev): the Emergent preview URL — where new changes land before deploy.", "p"),
    ("Always record which environment a defect was found on.", "p"),
    ("", ""),
    ("Test accounts", "h"),
    ("• Admin:    admin@bigfm.dubai / AdminPass123!", "p"),
    ("• Merchant: test@bigfm.dubai  / TestUser123!", "p"),
    ("• User:     demo@bigfm.dubai  / DemoPass123!", "p"),
    ("", ""),
    ("How to use this workbook", "h"),
    ("1. Go to the 'UAT Test Cases' tab.", "p"),
    ("2. Each row is one test. Follow 'Test Steps' with the given 'Test Data'.", "p"),
    ("3. Compare what you see to 'Expected Result'.", "p"),
    ("4. Fill 'Actual Result', set 'Status' (Pass / Fail / Blocked / N/A) from the dropdown,", "p"),
    ("   add your name in 'Tester', the 'Date', and any 'Notes / Defect ID'.", "p"),
    ("5. The 'Summary' tab auto-counts Pass/Fail/Blocked as you fill Status.", "p"),
    ("", ""),
    ("Priority key", "h"),
    ("• P0 = critical (must pass to ship)   • P1 = high   • P2 = medium/nice-to-have", "p"),
    ("Platform: Web = radiolife.app / preview web · Mobile = Expo app · Both = verify on each.", "p"),
    ("", ""),
    ("Tip: use Data > Filter on the test-cases tab to run by Module, Priority, or Platform.", "p"),
]
r = 1
for text, kind in lines:
    c = ws0.cell(row=r, column=2, value=text)
    if kind == "title":
        c.font = Font(bold=True, size=18, color=ACCENT)
    elif kind == "h":
        c.font = Font(bold=True, size=13, color=BRAND)
    else:
        c.font = Font(size=11)
    r += 1

# ---------- Sheet 2: UAT Test Cases ----------
ws = wb.create_sheet("UAT Test Cases")
ws.sheet_view.showGridLines = False
ws.append(HEADERS)
for col, _ in enumerate(HEADERS, start=1):
    cell = ws.cell(row=1, column=col)
    cell.fill = HDR_FILL
    cell.font = HDR_FONT
    cell.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
    cell.border = BORDER

# group counter per module for nice IDs
mod_abbr = {}
def _abbr(m):
    if m not in mod_abbr:
        mod_abbr[m] = "".join(w[0] for w in m.replace("/", " ").split()[:2]).upper()
    return mod_abbr[m]
counters = {}

for (module, feature, prio, plat, typ, pre, steps, data, exp) in CASES:
    ab = _abbr(module)
    counters[ab] = counters.get(ab, 0) + 1
    tc_id = f"{ab}-{counters[ab]:02d}"
    row = [tc_id, module, feature, prio, plat, typ, pre, steps, data, exp, "", "", "", "", ""]
    ws.append(row)
    ridx = ws.max_row
    for col in range(1, len(HEADERS) + 1):
        cell = ws.cell(row=ridx, column=col)
        cell.alignment = WRAP
        cell.border = BORDER
        cell.font = Font(size=10)
    # priority colour
    pc = ws.cell(row=ridx, column=4)
    pc.alignment = CENTER
    if prio in PRIO_FILL:
        pc.fill = PatternFill("solid", fgColor=PRIO_FILL[prio])
    ws.cell(row=ridx, column=5).alignment = CENTER  # platform
    ws.cell(row=ridx, column=6).alignment = CENTER  # type

# widths
widths = [10, 20, 26, 8, 9, 12, 30, 40, 22, 46, 30, 12, 12, 12, 22]
for i, w in enumerate(widths, start=1):
    ws.column_dimensions[get_column_letter(i)].width = w
ws.freeze_panes = "A2"
ws.auto_filter.ref = f"A1:{get_column_letter(len(HEADERS))}{ws.max_row}"

# Status dropdown
dv = DataValidation(type="list", formula1='"Pass,Fail,Blocked,N/A,Not Run"', allow_blank=True)
ws.add_data_validation(dv)
dv.add(f"L2:L{ws.max_row}")
for rr in range(2, ws.max_row + 1):
    ws.cell(row=rr, column=12, value="Not Run").alignment = CENTER

LAST = ws.max_row

# ---------- Sheet 3: Summary ----------
sm = wb.create_sheet("Summary")
sm.sheet_view.showGridLines = False
sm.column_dimensions["A"].width = 26
for col in "BCDEF":
    sm.column_dimensions[col].width = 14
sm["A1"] = "RadioLife UAT — Summary"
sm["A1"].font = Font(bold=True, size=16, color=ACCENT)

sm["A3"] = "Overall status"
sm["A3"].font = Font(bold=True, size=12)
tab = "'UAT Test Cases'"
stat_rows = [
    ("Total test cases", f"=COUNTA({tab}!A2:A{LAST})"),
    ("Pass", f'=COUNTIF({tab}!L2:L{LAST},"Pass")'),
    ("Fail", f'=COUNTIF({tab}!L2:L{LAST},"Fail")'),
    ("Blocked", f'=COUNTIF({tab}!L2:L{LAST},"Blocked")'),
    ("N/A", f'=COUNTIF({tab}!L2:L{LAST},"N/A")'),
    ("Not Run", f'=COUNTIF({tab}!L2:L{LAST},"Not Run")'),
    ("% Pass", f'=IF(COUNTA({tab}!A2:A{LAST})=0,0,COUNTIF({tab}!L2:L{LAST},"Pass")/COUNTA({tab}!A2:A{LAST}))'),
]
r = 4
for label, formula in stat_rows:
    sm.cell(row=r, column=1, value=label).font = Font(size=11)
    c = sm.cell(row=r, column=2, value=formula)
    c.font = Font(bold=True, size=11)
    if label == "% Pass":
        c.number_format = "0%"
    r += 1

# by-priority + by-module counts
sm.cell(row=4, column=4, value="By Priority").font = Font(bold=True)
pr = 5
for p in ["P0", "P1", "P2"]:
    sm.cell(row=pr, column=4, value=p)
    sm.cell(row=pr, column=5, value=f'=COUNTIF({tab}!D2:D{LAST},"{p}")').font = Font(bold=True)
    pr += 1

modules = []
for (m, *_rest) in CASES:
    if m not in modules:
        modules.append(m)
sm.cell(row=pr + 1, column=4, value="By Module (total / pass)").font = Font(bold=True)
mr = pr + 2
for m in modules:
    sm.cell(row=mr, column=4, value=m).font = Font(size=10)
    sm.cell(row=mr, column=5, value=f'=COUNTIF({tab}!B2:B{LAST},"{m}")')
    sm.cell(row=mr, column=6,
            value=f'=COUNTIFS({tab}!B2:B{LAST},"{m}",{tab}!L2:L{LAST},"Pass")')
    mr += 1
sm.cell(row=(pr + 1), column=6, value="").font = Font(size=10)

wb.save(OUT)
print("wrote", OUT, "| total cases:", len(CASES), "| modules:", len(modules))
