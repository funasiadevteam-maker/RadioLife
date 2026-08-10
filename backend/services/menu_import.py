"""AI + file menu import for RadioLife Marketplace merchants.

Parses a restaurant menu from an uploaded image, a PDF, or a website URL
using a vision-capable LLM (Gemini via the Emergent universal key), and
parses tabular product data from Excel/CSV. Returns a normalised list of
menu items: {section, name, description, price_cents}.
"""
from __future__ import annotations

import base64
import csv
import io
import json
import os
import re
import tempfile
from typing import Any, Dict, List

EMERGENT_LLM_KEY = os.environ.get("EMERGENT_LLM_KEY", "")
MENU_MODEL = ("gemini", "gemini-2.5-flash")

SYSTEM = (
    "You extract restaurant menus into structured JSON. Given a menu (image, "
    "PDF or website text), return ONLY a JSON object of the form "
    '{"items":[{"section":"","name":"","description":"","price":0.0}]}. '
    "section is the menu category (e.g. Starters, Mains, Drinks) — infer a "
    "sensible one if absent. price is a number in the menu's currency units "
    "(e.g. 12.5), 0 if unknown. Do not invent items. Return at most 60 items."
)


def _to_cents(price: Any) -> int:
    try:
        return max(0, int(round(float(price) * 100)))
    except (TypeError, ValueError):
        return 0


def _parse_json_items(text: str) -> List[Dict[str, Any]]:
    if not text:
        return []
    m = re.search(r"\{.*\}", text, re.DOTALL)
    raw = m.group(0) if m else text
    try:
        data = json.loads(raw)
    except json.JSONDecodeError:
        return []
    items = data.get("items", data) if isinstance(data, dict) else data
    out: List[Dict[str, Any]] = []
    if isinstance(items, list):
        for it in items:
            if not isinstance(it, dict):
                continue
            name = (it.get("name") or "").strip()
            if not name:
                continue
            out.append({
                "section": (it.get("section") or "Menu").strip()[:60],
                "name": name[:120],
                "description": (it.get("description") or "").strip()[:300],
                "price_cents": _to_cents(it.get("price")),
            })
    return out[:60]


async def _run_llm(user_text: str, *, image_b64: str | None = None,
                   file_path: str | None = None, mime: str | None = None) -> List[Dict[str, Any]]:
    from emergentintegrations.llm.chat import (
        LlmChat, UserMessage, ImageContent, FileContentWithMimeType)
    import uuid as _uuid

    chat = (LlmChat(api_key=EMERGENT_LLM_KEY, session_id=f"menu_{_uuid.uuid4().hex}",
                    system_message=SYSTEM)
            .with_model(*MENU_MODEL))
    file_contents = []
    if image_b64:
        file_contents.append(ImageContent(image_base64=image_b64))
    if file_path and mime:
        file_contents.append(FileContentWithMimeType(file_path=file_path, mime_type=mime))
    msg = UserMessage(text=user_text, file_contents=file_contents or None)
    resp = await chat.send_message(msg)
    text = resp if isinstance(resp, str) else getattr(resp, "content", str(resp))
    return _parse_json_items(text)


async def import_from_image(image_b64: str) -> List[Dict[str, Any]]:
    return await _run_llm("Extract every menu item from this menu image.",
                          image_b64=image_b64)


async def import_from_pdf(pdf_b64: str) -> List[Dict[str, Any]]:
    raw = base64.b64decode(pdf_b64)
    with tempfile.NamedTemporaryFile(suffix=".pdf", delete=False) as f:
        f.write(raw)
        path = f.name
    try:
        return await _run_llm("Extract every menu item from this menu PDF.",
                              file_path=path, mime="application/pdf")
    finally:
        try:
            os.unlink(path)
        except OSError:
            pass


async def import_from_url(url: str) -> List[Dict[str, Any]]:
    import httpx
    from bs4 import BeautifulSoup

    async with httpx.AsyncClient(timeout=25, follow_redirects=True,
                                 headers={"User-Agent": "Mozilla/5.0 RadioLifeBot"}) as client:
        r = await client.get(url)
        r.raise_for_status()
        html = r.text
    soup = BeautifulSoup(html, "html.parser")
    for tag in soup(["script", "style", "nav", "footer", "header"]):
        tag.decompose()
    text = re.sub(r"\n{2,}", "\n", soup.get_text("\n")).strip()[:16000]
    return await _run_llm(
        "The following is text scraped from a restaurant website. Extract the "
        f"menu items you can find.\n\n{text}")


def import_from_spreadsheet(file_b64: str, filename: str) -> List[Dict[str, Any]]:
    """Parse an .xlsx or .csv with columns like section/name/description/price."""
    raw = base64.b64decode(file_b64)
    rows: List[List[str]] = []
    if filename.lower().endswith(".csv"):
        text = raw.decode("utf-8", errors="ignore")
        rows = [r for r in csv.reader(io.StringIO(text))]
    else:
        from openpyxl import load_workbook
        wb = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
        ws = wb.active
        for row in ws.iter_rows(values_only=True):
            rows.append([("" if c is None else str(c)) for c in row])
    if not rows:
        return []
    header = [str(h).strip().lower() for h in rows[0]]

    def col(*names: str) -> int:
        for n in names:
            if n in header:
                return header.index(n)
        return -1

    ci_name = col("name", "item", "product", "title")
    ci_sec = col("section", "category", "group")
    ci_desc = col("description", "desc", "details")
    ci_price = col("price", "amount", "cost")
    # If no header match, assume positional: name, description, price, section.
    positional = ci_name < 0
    out: List[Dict[str, Any]] = []
    data_rows = rows if positional else rows[1:]
    for r in data_rows:
        def g(i: int) -> str:
            return r[i].strip() if 0 <= i < len(r) and r[i] is not None else ""
        if positional:
            name, desc, price, sec = (g(0), g(1), g(2), g(3))
        else:
            name, desc, price, sec = (g(ci_name), g(ci_desc), g(ci_price), g(ci_sec))
        if not name or name.lower() in ("name", "item", "product"):
            continue
        pnum = re.sub(r"[^0-9.]", "", price)
        out.append({
            "section": (sec or "Menu")[:60], "name": name[:120],
            "description": desc[:300], "price_cents": _to_cents(pnum) if pnum else 0,
        })
    return out[:200]
